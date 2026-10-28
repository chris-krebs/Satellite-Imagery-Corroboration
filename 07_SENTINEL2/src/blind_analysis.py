# /// script
# requires-python = ">=3.10"
# dependencies = ["openpyxl"]
# ///
"""Unblind the Sentinel-2 image labels and compare them with the truth, dNBR and FIRMS.

    uv run src/blind_analysis.py outputs/blind_review_filled.csv outputs/blind_review_2022.xlsx

Several label files (.csv or .xlsx) can be given; they are combined.

Run this only after every image is labelled. It reads the labels (any of the separators a
spreadsheet app may have used), joins blind ids to events, and writes
outputs/blind_results.csv plus a summary on screen.

Seen-before events (full panels looked at before labelling) are reported separately.
"""
import csv
import re
import sys
from collections import Counter, defaultdict
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
PROJECT = ROOT.parent
OUT = ROOT / "outputs"
SEEN_BEFORE = {"E006", "E017", "E021"}           # full panels were opened during the first test run

LABEL_FIX = {"not_assassable": "not_assessable", "not_assesable": "not_assessable", "new_damage_scar": "new_damage_or_scar",
             "new_damage": "new_damage_or_scar", "scar": "new_damage_or_scar", "fire": "active_fire"}
CONF_FIX = {"meidum": "medium", "med": "medium", "hi": "high"}
LABELS = {"active_fire", "smoke", "new_damage_or_scar", "nothing_visible", "not_assessable"}
POSITIVE = {"active_fire", "smoke", "new_damage_or_scar"}


def lines_of(path):
    """Text lines of a .csv, or of an .xlsx turned into ';'-separated lines."""
    path = Path(path)
    if path.suffix.lower() == ".xlsx":
        from openpyxl import load_workbook
        ws = load_workbook(path, data_only=True).worksheets[0]
        return [";".join("" if v is None else str(v) for v in row) for row in ws.iter_rows(values_only=True)]
    return path.read_text(encoding="utf-8-sig").splitlines()


def read_labels(path):
    """blind_id -> (label, confidence). Tolerates mixed ',' / ';' and stray header lines."""
    out, problems = {}, []
    for line in lines_of(path):
        parts = [p.strip() for p in re.split(r"[;,\t]", line)]
        bid = parts[0] if parts else ""
        if not re.fullmatch(r"B[0-9A-F]{6}", bid):
            continue
        words = [p.lower() for p in parts[1:] if p and not p.endswith(".png")]
        label = next((LABEL_FIX.get(w, w) for w in words if LABEL_FIX.get(w, w) in LABELS), "")
        conf = next((CONF_FIX.get(w, w) for w in words if CONF_FIX.get(w, w) in ("high", "medium", "low")), "")
        if not label:
            problems.append(f"{bid}: no recognisable label in {words}")
        out[bid] = (label, conf)
    return out, problems


def read_csv(path):
    if not Path(path).exists():
        return []
    return list(csv.DictReader(open(path, encoding="utf-8")))


def main():
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    labels, problems = {}, []
    for path in sys.argv[1:]:
        got, probs = read_labels(path)
        dup = set(got) & set(labels)
        if dup:
            problems.append(f"{path}: labelled again, later file wins: {', '.join(sorted(dup))}")
        labels.update(got)
        problems += probs
    key = {r["blind_id"]: r["event_id"] for r in read_csv(OUT / "blind_key.csv")}
    s2 = {r["event_id"]: r for r in read_csv(OUT / "s2_summary.csv")}
    ev = {r["event_id"]: r for r in read_csv(PROJECT / "02_OSM_CHECK" / "data" / "events_reviewed.csv")}
    heat = {r["event_id"]: r for r in read_csv(PROJECT / "05_PERSISTENT_HEAT" / "outputs" / "heat_event_summary.csv")}

    missing = sorted(set(key) - {b for b, (lab, _) in labels.items() if lab})
    rows = []
    for bid, (label, conf) in sorted(labels.items()):
        if not label:
            continue                         # reported above; not part of any count
        eid = key.get(bid)
        if not eid:
            problems.append(f"{bid}: not in blind_key.csv")
            continue
        e, s, h = ev.get(eid, {}), s2.get(eid, {}), heat.get(eid, {})
        rows.append({
            "blind_id": bid, "event_id": eid, "category": e.get("category", ""), "place_name": e.get("place_name", ""),
            "is_control": e.get("category") == "negative_control", "seen_before": eid in SEEN_BEFORE,
            "label": label, "confidence": conf, "visual_positive": label in POSITIVE,
            "s2_status": s.get("status", ""), "first_after_days": s.get("first_after_days", ""),
            "first_after_cloud": s.get("first_after_cloud", ""), "clear_after_days": s.get("clear_after_days", ""),
            "dnbr_site_minus_ring": s.get("dnbr_site_minus_ring", ""),
            "frac_site_dnbr_gt_0_27": s.get("frac_site_dnbr_gt_0_27", ""),
            "area_site_dnbr_gt_0_27_m2": s.get("area_site_dnbr_gt_0_27_m2", ""),
            "excess_frac_gt_0_27": s.get("excess_frac_gt_0_27", ""),
            "firms_site_signal": h.get("n_site_signal", ""), "firms_site_raw": h.get("n_site_raw", ""),
        })
    with open(OUT / "blind_results.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0].keys())); w.writeheader(); w.writerows(rows)

    for p in problems:
        print("note:", p)
    if missing:
        print("not labelled:", ", ".join(missing))
    blind = [r for r in rows if not r["seen_before"]]
    print(f"\n{len(rows)} labelled images ({len(rows) - len(blind)} seen before labelling, reported separately)\n")
    def table(rs, title):
        c = defaultdict(Counter)
        for r in rs:
            c["control" if r["is_control"] else "event"][r["label"]] += 1
        print(title)
        order = ["active_fire", "smoke", "new_damage_or_scar", "nothing_visible", "not_assessable"]
        print(f"  {'':10}" + "".join(f"{o[:18]:>20}" for o in order) + f"{'total':>8}")
        for k in ("event", "control"):
            print(f"  {k:10}" + "".join(f"{c[k][o]:>20}" for o in order) + f"{sum(c[k].values()):>8}")
    table(blind, "Blind images: label by truth")
    def rate(rs):
        assess = [r for r in rs if r["label"] != "not_assessable"]
        pos = [r for r in assess if r["visual_positive"]]
        return len(pos), len(assess), len(rs)
    for name, sel in (("events", [r for r in blind if not r["is_control"]]), ("controls", [r for r in blind if r["is_control"]])):
        p, a, n = rate(sel)
        print(f"  {name}: {p} of {a} assessable images called fire/smoke/damage ({n - a} not assessable)")
    print("\nBy category (blind, events only): visual positive / assessable")
    by = defaultdict(list)
    for r in blind:
        if not r["is_control"]:
            by[r["category"]].append(r)
    for cat, rs in sorted(by.items()):
        p, a, n = rate(rs)
        print(f"  {cat:22} {p} / {a}   (+{n - a} not assessable)")
    print("\nSeen-before images:", ", ".join(f"{r['event_id']}={r['label']}" for r in rows if r["seen_before"]))
    print(f"\nwritten: {OUT / 'blind_results.csv'}")


if __name__ == "__main__":
    main()
